"""The app boots, serves its pages, and reconciles a synthetic statement and ledger end to end (monthly and
daily mode), including the Excel exports. Inputs come from engine/simulacion.py (the demo's simulated bank
and FBS) for a fixed past range, so they are always the same."""
import io
import time
from datetime import date, datetime

import openpyxl
from conftest import SEED_ACCOUNT

from engine import simulacion

ACCOUNT_ID = "ciudad-350123456789-ars"
ACCOUNT = {**SEED_ACCOUNT, "cuenta_id": ACCOUNT_ID, "plan_e": 1141121, "plan_o": 1141122, "rol": "principal"}
FROM, TO = date(2026, 9, 14), date(2026, 9, 15)
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _statement() -> tuple[str, bytes, int]:
    """The account's bank statement for the range (Banco Ciudad .csv) and its number of movements."""
    return simulacion.extracto(ACCOUNT, FROM, TO)


def _fbs_ledger_xlsx() -> bytes:
    """An FBS "Consulta del Mayor" export (sheets E and O) with the simulated entries of the account."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, code in (("E", ACCOUNT["plan_e"]), ("O", ACCOUNT["plan_o"])):
        ws = wb.create_sheet(f"({sheet}) ({code})")
        ws.append(["Consulta del Mayor"])
        ws.append([f"Cuenta: CIUDAD CC $ {ACCOUNT['numero']} ({sheet}) ({code})"])
        ws.append([])
        ws.append(["Asiento", "Fecha", "Referencia", "Comentario", "Debe", "Haber"])
        for e in simulacion.asientos(ACCOUNT, FROM, TO):
            if e.hoja == sheet:
                ws.append([e.numero, datetime(e.fecha.year, e.fecha.month, e.fecha.day), e.referencia,
                           e.comentario, e.debe, e.haber])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _wait_for_job(client, job_id):
    for _ in range(400):
        progress = client.get(f"/api/progreso/{job_id}").json()
        if progress["estado"] != "procesando":
            return progress
        time.sleep(0.05)
    raise AssertionError("the reconciliation job did not finish")


def test_health(client):
    r = client.get("/health")
    assert r.status_code == 200
    assert r.json()["status"] == "ok"


def test_pages_are_served(client):
    for path in ("/", "/diario"):
        r = client.get(path)
        assert r.status_code == 200
        assert "<html" in r.text.lower()


def test_monthly_reconciliation_and_export(client):
    name, statement, movements = _statement()
    r = client.post("/api/conciliar", files=[("extractos", (name, statement)),
                                             ("mayor", ("mayor.xlsx", _fbs_ledger_xlsx()))])
    assert r.status_code == 200, r.text
    job_id = r.json()["job_id"]
    assert _wait_for_job(client, job_id)["estado"] == "listo"

    summary = client.get(f"/api/resultado/{job_id}").json()["resumen"]
    assert summary["movimientos_banco"] == movements == 49
    assert summary["conciliados_e"] == 35
    assert summary["e_sin_banco"]["cantidad"] == 2

    r = client.get(f"/api/exportar/{job_id}")
    assert r.status_code == 200
    assert r.headers["content-type"] == XLSX
    assert "Resumen" in openpyxl.load_workbook(io.BytesIO(r.content)).sheetnames


def test_daily_reconciliation_history_and_delete(client):
    name, statement, movements = _statement()
    r = client.post("/api/diario/identificar",
                    files=[("archivos", (name, statement)),
                           ("archivos", ("mayor-fbs.xlsx", _fbs_ledger_xlsx()))])
    assert r.status_code == 200, r.text
    staging = r.json()
    assert sorted(f["tipo"] for f in staging["archivos"]) == ["extracto", "fbs"]

    assignments = {f["archivo"]: ACCOUNT_ID for f in staging["archivos"]}
    r = client.post(f"/api/diario/conciliar/{staging['staging_id']}", json={"asignaciones": assignments})
    assert r.status_code == 200, r.text
    board = r.json()
    [account] = board["cuentas"]
    assert account["estado"] == "ok"
    assert account["resumen"]["movimientos_banco"] == movements
    assert account["resumen"]["conciliados"] == 38

    history = client.get("/api/diario").json()["grupos"]
    assert [g["grupo_id"] for g in history] == [board["grupo_id"]]
    assert client.get(f"/api/diario/{board['grupo_id']}").status_code == 200
    r = client.get(f"/api/diario/{board['grupo_id']}/excel")
    assert r.status_code == 200 and r.headers["content-type"] == XLSX

    assert client.delete(f"/api/diario/{board['grupo_id']}").status_code == 200
    assert client.get("/api/diario").json()["grupos"] == []
    assert client.get(f"/api/resultado/{account['job_id']}").status_code == 404
